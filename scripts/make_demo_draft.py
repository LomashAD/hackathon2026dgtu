#!/usr/bin/env python
"""Создаёт samples/Черновик_с_ошибками.docx — черновик распоряжения с типичными ошибками для демо проверки."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx  # noqa: E402
from docx.shared import Mm, Pt  # noqa: E402

from app.config import ROOT  # noqa: E402

SRC = ROOT / "data" / "Модуль_распоряжения_о_создании_рабочей_группы.docx"
OUT = ROOT / "samples" / "Черновик_с_ошибками.docx"


def main() -> None:
    d = docx.Document(str(SRC))
    body = d.element.body
    started = False
    for el in list(body.iterchildren()):
        text = "".join(el.itertext()).strip()
        if text.startswith("Примечания по заполнению"):
            started = True
        if started and el.tag.endswith(("}p", "}tbl")):
            body.remove(el)
    for p in d.paragraphs:
        t = p.text.strip()
        repl = None
        if t.startswith("О создании рабочей группы по"):
            repl = "О создании рабочей группы по подготовке Регламента электронного архива"
        elif t.startswith("В целях"):
            repl = "В целях подготовки проекта Регламента электронного архива [1],"
        elif t == "ОБЯЗЫВАЮ:":
            repl = "Обязываю"
        elif t.startswith("1. Создать рабочую группу"):
            repl = "1. Создать рабочую группу по подготовке Регламента электронного архива (далее - Рабочая группа) в составе:"
        elif t.startswith("2. Рабочей группе до"):
            repl = "2. Рабочей группе до 15/12/2026 разработать проект Регламента и представить ____________________."
        elif t.startswith("5. Рабочая группа прекращает"):
            repl = "5. Рабочая группа прекращает деятельность после утверждения Регламента."
        elif t.startswith("6. Контроль"):
            p._element.getparent().remove(p._element)
            continue
        if repl is not None:
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            p.runs[0].text = repl
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(12)
    for t in d.tables:
        for row in t.rows:
            if row.cells[0].text.strip() == "Руководитель":
                row.cells[1].paragraphs[0].runs[0].text = "начальник отдела ДОУ И. И. Иванов"
    d.sections[0].left_margin = Mm(20)
    OUT.parent.mkdir(exist_ok=True)
    d.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
