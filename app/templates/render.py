"""Генерация .docx: исходный модуль (data/) → docxtpl-шаблон (теги Jinja вместо линий) → черновик.

Сборка шаблона выполняется кодом по секции layout/tables YAML-описания, оформление исходного
модуля (шрифты, выравнивание, служебные зоны «НЕ УДАЛЯТЬ») сохраняется. По методическим
рекомендациям (п. 3.3) удаляются гриф «УТВЕРЖДАЮ», пометка «МОДУЛЬНЫЙ ДОКУМЕНТ» и примечания.
"""
from __future__ import annotations

import copy
import datetime as dt
import re
import secrets
import time
from pathlib import Path

from app.config import ROOT, get_config, path


def _set_text(paragraph, text: str) -> None:
    """Заменить текст абзаца, сохранив формат первого прогона."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _delete_paragraph(paragraph) -> None:
    el = paragraph._element
    el.getparent().remove(el)


def _iter_cells(doc):
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield t, row, c


def build_template(tpl: dict, out_path: Path) -> Path:
    import docx

    src = ROOT / tpl["source_file"]
    d = docx.Document(str(src))
    body = d.element.body

    # 1. Удаляем «Примечания по заполнению…» и всё после них.
    started = False
    for el in list(body.iterchildren()):
        tag = el.tag.split("}")[1]
        text = "".join(el.itertext()).strip()
        if tag == "p" and text.startswith("Примечания по заполнению"):
            started = True
        if started and tag in ("p", "tbl"):
            body.remove(el)

    # 2. Удаляем гриф УТВЕРЖДАЮ (первая таблица) и пометку МОДУЛЬНЫЙ ДОКУМЕНТ.
    for t in list(d.tables):
        if "УТВЕРЖДАЮ" in "".join(t._tbl.itertext()):
            t._tbl.getparent().remove(t._tbl)
    for p in list(d.paragraphs):
        if p.text.strip() == "МОДУЛЬНЫЙ ДОКУМЕНТ":
            _delete_paragraph(p)

    # 3. Абзацы с полями → теги Jinja.
    used = set()
    for rule in tpl.get("layout", []):
        for p in d.paragraphs:
            if id(p._p) in used:
                continue
            if p.text.strip().startswith(rule["match"]):
                _set_text(p, rule["text"])
                used.add(id(p._p))
                if rule.get("if"):
                    before = p.insert_paragraph_before("{%p if " + rule["if"] + " %}")
                    nxt = p._p.getnext()
                    after = before.__class__(copy.deepcopy(before._p), p._parent)
                    _set_text(after, "{%p endif %}")
                    if nxt is not None:
                        nxt.addprevious(after._p)
                    else:
                        p._p.addnext(after._p)
                break
        else:
            raise ValueError(f"{tpl['id']}: не найден абзац, начинающийся с {rule['match']!r}")

    # 4. Ячейки таблиц.
    for rule in tpl.get("tables", []):
        done = False
        for _t, row, _c in _iter_cells(d):
            cells = row.cells
            if cells[0].text.strip().startswith(rule["match"]) or (
                    rule["match"] in "".join(c.text for c in cells) and rule["match"] == "И. О. Фамилия"):
                target = cells[rule["cell"]]
                ps = target.paragraphs
                _set_text(ps[0], rule["text"])
                for extra in ps[1:]:
                    _delete_paragraph(extra)
                done = True
                break
        if not done:
            raise ValueError(f"{tpl['id']}: не найдена строка таблицы {rule['match']!r}")

    # 5. Таблицы с переменным числом строк (подразделение — ответственный).
    for rule in tpl.get("row_tables", []):
        for t in d.tables:
            if t.rows[0].cells[0].text.strip().startswith(rule["header"]):
                rows = t.rows[1:]
                proto = rows[0]._tr
                for r in rows[1:]:
                    t._tbl.remove(r._tr)
                loop_open = copy.deepcopy(proto)
                loop_close = copy.deepcopy(proto)
                proto.addprevious(loop_open)
                proto.addnext(loop_close)
                from docx.table import _Row
                r_open, r_body, r_close = (_Row(x, t) for x in (loop_open, proto, loop_close))
                _set_text(r_open.cells[0].paragraphs[0], "{%tr for r in " + rule["var"] + " %}")
                _set_text(r_close.cells[0].paragraphs[0], "{%tr endfor %}")
                for i, col in enumerate(rule["cols"]):
                    _set_text(r_body.cells[i].paragraphs[0], "{{ r." + col + " }}")
                break

    out_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_path))
    return out_path


def template_docx(tpl: dict, rebuild: bool = False) -> Path:
    out = path("templates_docx") / f"{tpl['id']}.docx"
    src = ROOT / tpl["source_file"]
    if rebuild or not out.exists() or out.stat().st_mtime < src.stat().st_mtime \
            or out.stat().st_mtime < Path(tpl["_yaml_path"]).stat().st_mtime:
        build_template(tpl, out)
    return out


OWNER_SUFFIX = ".owner"   # рядом с черновиком — логин автора: скачать файл через API может только он


def draft_owner(file: Path) -> str:
    f = file.with_name(file.name + OWNER_SUFFIX)
    return f.read_text(encoding="utf-8").strip() if f.exists() else ""


def purge_drafts(out_dir: Path | None = None, ttl_hours: float | None = None) -> int:
    """Удалить черновики старше security.drafts_ttl_hours (в них ФИО и должности сотрудников)."""
    out_dir = out_dir or path("output_dir")
    ttl = ttl_hours if ttl_hours is not None else get_config()["security"].get("drafts_ttl_hours", 24)
    if not ttl or not out_dir.exists():
        return 0
    old = time.time() - ttl * 3600
    n = 0
    for f in out_dir.iterdir():
        if f.is_file() and f.stat().st_mtime < old:
            f.unlink(missing_ok=True)
            n += f.suffix == ".docx"
    return n


def render(tpl: dict, context: dict, out_dir: Path | None = None, owner: str = "") -> Path:
    from docxtpl import DocxTemplate

    tp = DocxTemplate(str(template_docx(tpl)))
    tp.render(context, autoescape=True)
    out_dir = out_dir or path("output_dir")
    out_dir.mkdir(parents=True, exist_ok=True)
    purge_drafts(out_dir)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    name = re.sub(r"[^\w]+", "_", tpl["title"])[:60]
    # случайная часть: имя черновика не подобрать по шаблону и времени
    out = out_dir / f"Черновик_{name}_{stamp}_{secrets.token_hex(4)}.docx"
    tp.save(str(out))
    if owner:
        out.with_name(out.name + OWNER_SUFFIX).write_text(owner, encoding="utf-8")
    return out
