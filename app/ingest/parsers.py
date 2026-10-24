"""Парсеры исходных документов в структуру «глава → пункт → подпункт → приложение».

Каждый парсер возвращает список секций:
    {"section_path": "5.4.1", "chapter": "5", "chapter_title": "...",
     "appendix": "" | "1", "text": "...", "kind": "point" | "appendix" | "preamble" | "fragment"}
"""
from __future__ import annotations

import re
from pathlib import Path

POINT_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){1,3})\.\s+(.*)$")
CHAPTER_RE = re.compile(r"^Глава\s+(\d+)\.\s*(.*)$")
APPENDIX_RE = re.compile(r"^Приложение\s*№\s*(\d+)\s*$")


def _is_watermark(obj) -> bool:
    # Водяной знак «УЧЕБНЫЙ ПРИМЕР...» набран повёрнутыми символами.
    return obj.get("object_type") == "char" and abs(obj["matrix"][1]) > 0.01


def _join(prev: str, line: str) -> str:
    if not prev:
        return line
    if prev.endswith("-") and not prev.endswith(" -"):
        return prev + line  # перенос составного слова: «официально-» + «делового»
    return prev + " " + line


def _pdf_lines(pdf_path: Path):
    """Строки PDF без водяного знака: (text, x0, is_bold, page_no)."""
    import pdfplumber

    with pdfplumber.open(str(pdf_path)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            clean = page.filter(lambda o: not _is_watermark(o))
            for ln in clean.extract_text_lines():
                text = ln["text"].strip()
                if not text:
                    continue
                bold = any("Bold" in c.get("fontname", "") for c in ln["chars"][:3])
                yield text, ln["x0"], bold, page_no


def _key(section_id: str) -> tuple:
    return tuple(int(x) for x in section_id.split("."))


def _follows(prev_id: str, new_id: str) -> bool:
    """Номер пункта должен идти строго после предыдущего (защита от ложных совпадений)."""
    return not prev_id or _key(new_id) > _key(prev_id)


def parse_instruction_pdf(pdf_path: Path, para_indent: float = 100.0) -> list[dict]:
    """Инструкция: главы «Глава N.», пункты «N.M.», подпункты «N.M.K.», «Приложение № N».

    Абзацы восстанавливаются по отступу первой строки (x0 > para_indent) и маркеру «•».
    """
    sections: list[dict] = []
    chapter, chapter_title = "", ""
    current: dict | None = None
    in_heading = False
    title_lines: list[str] = []
    last_id = ""

    def close():
        nonlocal current
        if current:
            current["text"] = "\n".join(p.strip() for p in current["paras"] if p.strip())
            del current["paras"]
            sections.append(current)
        current = None

    for text, x0, bold, page in _pdf_lines(pdf_path):
        m_ch = CHAPTER_RE.match(text)
        if m_ch and bold:
            close()
            chapter, chapter_title = m_ch.group(1), m_ch.group(2).strip()
            in_heading = True
            continue
        if in_heading and bold:
            chapter_title = (chapter_title + " " + text).strip()
            continue
        in_heading = False

        m_app = APPENDIX_RE.match(text)
        if m_app:
            close()
            current = {"section_path": f"Приложение {m_app.group(1)}", "chapter": "",
                       "chapter_title": "", "appendix": m_app.group(1), "kind": "appendix",
                       "page": page, "paras": []}
            title_lines = []
            continue

        m_pt = POINT_RE.match(text)
        if m_pt and chapter and _follows(last_id, m_pt.group(1)):
            last_id = m_pt.group(1)
            close()
            current = {"section_path": m_pt.group(1), "chapter": chapter,
                       "chapter_title": chapter_title, "appendix": "", "kind": "point",
                       "page": page, "paras": [m_pt.group(2)]}
            continue

        if current is None:
            # Титульный лист (до главы 1) — служебные сведения об утверждении.
            if not sections and not chapter:
                title_lines.append(text)
            continue

        prev = current["paras"][-1] if current["paras"] else ""
        unfinished = bool(prev) and not re.search(r"[.;:!?»]$", prev)
        if text.startswith("•") or (x0 > para_indent and not unfinished) or (
            current["kind"] == "appendix" and re.match(r"^\d+\.\s", text)
        ):
            current["paras"].append(text)
        else:
            current["paras"][-1] = _join(current["paras"][-1], text)
    close()
    return sections


def parse_title_pdf(pdf_path: Path) -> str:
    """Первые строки PDF (гриф утверждения + название)."""
    lines = []
    for text, _x0, _bold, page in _pdf_lines(pdf_path):
        if page > 1:
            break
        lines.append(text)
    return " ".join(lines)


def _docx_paragraphs(docx_path: Path) -> list[str]:
    import docx

    d = docx.Document(str(docx_path))
    return [p.text.strip() for p in d.paragraphs if p.text.strip()]


def parse_numbered_docx(docx_path: Path) -> list[dict]:
    """DOCX с разделами «N. Заголовок» и пунктами «N.M. текст» (методические рекомендации)."""
    sections: list[dict] = []
    chapter, chapter_title = "", ""
    current: dict | None = None
    head_re = re.compile(r"^(\d{1,2})\.\s+(\S.*)$")
    for text in _docx_paragraphs(docx_path):
        m_pt = POINT_RE.match(text)
        m_head = head_re.match(text)
        if m_pt:
            current = {"section_path": m_pt.group(1), "chapter": chapter,
                       "chapter_title": chapter_title, "appendix": "", "kind": "point",
                       "text": m_pt.group(2)}
            sections.append(current)
        elif m_head and (current is None or len(text) < 80) and not text.endswith(";"):
            # Заголовок раздела («2. Создание регистрационной карточки проекта»);
            # длинные строки «1. ответственный ...;» — элементы перечня внутри пункта.
            if current is not None and current["chapter"] and current["text"].rstrip().endswith(":"):
                current["text"] += "\n" + text
                continue
            chapter, chapter_title = m_head.group(1), m_head.group(2)
        elif current is not None:
            current["text"] += "\n" + text
    return sections


def parse_docx_title(docx_path: Path) -> str:
    paras = _docx_paragraphs(docx_path)
    return " ".join(paras[:2])


# ---------------------------------------------------------------- матрица визирования

MATRIX_COLUMNS = ["legal", "dou", "hr", "other"]
MATRIX_COLUMN_TITLES = {
    "legal": "Юридическое подразделение",
    "dou": "Департамент ДОУ",
    "hr": "Кадровое подразделение",
}


def _cell_value(text: str) -> dict:
    lines = [" ".join(x.split()) for x in text.split("\n") if x.strip()]
    t = " ".join(lines)
    if not t or t.lower() == "нет":
        return {"required": "no"}
    if t.lower() == "да":
        return {"required": "yes"}
    if "если" in t.lower():
        unit, _, cond = t.partition(" визирует проект, ")
        return {"required": "conditional", "units": [unit.strip()], "condition": cond.strip().rstrip(".")}
    return {"required": "yes", "units": [u.strip() for ln in lines for u in ln.split(" / ") if u.strip()]}


def parse_matrix_docx(docx_path: Path) -> dict:
    """Матрица обязательного визирования → {"preamble": [...], "rows": [...]}."""
    import docx

    d = docx.Document(str(docx_path))
    preamble = [p.text.strip() for p in d.paragraphs if re.match(r"^\d\.\s", p.text.strip())]
    notes = [p.text.strip() for p in d.paragraphs if p.text.strip().startswith("Примечание")]
    title = " ".join(p.text.strip() for p in d.paragraphs[:2] if p.text.strip())
    rows = []
    group = ""
    table = d.tables[0]
    for r in table.rows[1:]:
        cells = []
        for c in r.cells:  # объединённые ячейки повторяются — схлопываем подряд идущие
            if not cells or c._tc is not cells[-1]._tc:
                cells.append(c)
        texts = [c.text.strip() for c in cells]
        if len(set(texts)) == 1:
            group = "internal" if "ВНУТРЕН" in texts[0].upper() else "external"
            continue
        m = re.match(r"^(\d+)\.\s*(.*)$", texts[0], re.S)
        num, name = (m.group(1), m.group(2)) if m else ("", texts[0])
        row = {"id": f"{group}-{num}", "group": group, "num": num, "title": " ".join(name.split())}
        for key, txt in zip(MATRIX_COLUMNS, texts[1:]):
            row[key] = _cell_value(txt)
        rows.append(row)
    return {"title": title, "preamble": preamble, "notes": notes, "rows": rows}


def matrix_sections(matrix: dict) -> list[dict]:
    """Текстовое представление матрицы для полнотекстового поиска (преамбула + строки)."""
    out = []
    for p in matrix["preamble"]:
        num, _, text = p.partition(". ")
        out.append({"section_path": num, "chapter": "", "chapter_title": "Общие положения",
                    "appendix": "", "kind": "point", "text": text})
    for row in matrix["rows"]:
        parts = [row["title"]]
        for key in MATRIX_COLUMNS:
            v = row[key]
            label = MATRIX_COLUMN_TITLES.get(key, "Прочие департаменты")
            if v["required"] == "yes":
                parts.append(f"{label}: {', '.join(v.get('units', [])) or 'Да'}")
            elif v["required"] == "conditional":
                parts.append(f"{label}: {', '.join(v['units'])} — {v['condition']}")
        grp = "Внутренние документы" if row["group"] == "internal" else "Внешние документы"
        out.append({"section_path": f"{grp}, строка {row['num']}", "chapter": "",
                    "chapter_title": grp, "appendix": "", "kind": "matrix_row",
                    "text": ". ".join(parts)})
    return out


# ---------------------------------------------------------------- произвольный документ

PLAIN_EXT = (".docx", ".pdf", ".txt", ".md")
_HEADING_STYLE = re.compile(r"^(heading|заголовок|title|название)", re.I)
_FRAGMENT_CHARS = 900


def _plain_blocks(path: Path) -> list[tuple[str, bool]]:
    """Абзацы документа без структуры «N.M.»: (текст, заголовок ли)."""
    ext = path.suffix.lower()
    out: list[tuple[str, bool]] = []
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        for para in d.paragraphs:
            t = para.text.strip()
            if t:
                style = para.style.name if para.style is not None else ""
                out.append((t, bool(_HEADING_STYLE.match(style)) and len(t) < 150))
        for table in d.tables:
            for row in table.rows:
                cells = list(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip()))
                if cells:
                    out.append((" | ".join(cells), False))
    elif ext == ".pdf":
        cur = ""
        for text, _x0, bold, _page in _pdf_lines(path):
            if bold and len(text) < 120:
                if cur:
                    out.append((cur, False))
                out.append((text, True))
                cur = ""
            elif cur and (re.match(r"^(\d+[.)]|[•–—-])\s", text) or (cur[-1] in ".:;!?" and text[:1].isupper())):
                out.append((cur, False))
                cur = text
            else:
                cur = _join(cur, text)
        if cur:
            out.append((cur, False))
    elif ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace")
        for block in re.split(r"\n\s*\n", text):
            lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
            for ln in [x for x in lines if x.startswith("#")]:
                out.append((ln.lstrip("#").strip(), True))
            body = " ".join(x for x in lines if not x.startswith("#"))
            if body:
                out.append((body, False))
    else:
        raise ValueError(f"Формат {ext or 'без расширения'} не поддерживается: нужен {', '.join(PLAIN_EXT)}")
    return out


def parse_plain(path: Path) -> list[dict]:
    """Любой документ: пункты «N.M.» — как пункты, остальной текст — фрагменты до ~900 символов
    в пределах раздела (заголовка). Для регламентов и памяток без строгой нумерации."""
    sections: list[dict] = []
    heading, buf, n, point = "", [], 0, None

    def flush() -> None:
        nonlocal buf, n
        if buf:
            n += 1
            sections.append({"section_path": str(n), "chapter": "", "chapter_title": heading, "appendix": "",
                             "kind": "fragment", "text": "\n".join(buf)})
            buf = []

    for text, is_heading in _plain_blocks(path):
        m = POINT_RE.match(text)
        if is_heading:
            flush()
            heading, point = text, None
        elif m and not any(s["section_path"] == m.group(1) for s in sections):
            flush()
            point = {"section_path": m.group(1), "chapter": "", "chapter_title": heading, "appendix": "",
                     "kind": "point", "text": m.group(2)}
            sections.append(point)
        elif point is not None:
            point["text"] += "\n" + text
        else:
            if buf and sum(len(x) for x in buf) + len(text) > _FRAGMENT_CHARS:
                flush()
            buf.append(text)
    flush()
    return sections


def parse_auto(path: Path) -> list[dict]:
    """Структурный парсер по расширению, если он охватил почти весь текст; иначе — parse_plain."""
    structured = {".pdf": parse_instruction_pdf, ".docx": parse_numbered_docx}.get(path.suffix.lower())
    plain = parse_plain(path)
    if structured:
        try:
            secs = structured(path)
        except Exception:  # noqa: BLE001 — не та структура
            secs = []
        total = sum(len(s["text"]) for s in plain) or 1
        if sum(s["kind"] == "point" for s in secs) >= 5 and sum(len(s["text"]) for s in secs) >= 0.8 * total:
            return secs
    return plain


PARSERS = {
    "instruction_pdf": parse_instruction_pdf,
    "numbered_docx": parse_numbered_docx,
    "plain": parse_plain,       # любой .docx/.pdf/.txt/.md: пункты «N.M.» и фрагменты по разделам
    "auto": parse_auto,         # выбор по содержимому — так подключаются документы из админ-панели
}
