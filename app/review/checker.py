"""Проверка загруженного черновика (.docx/.pdf) на соответствие Инструкции и модулю.

Все проверки детерминированные (правила кода), у каждой — ссылка на пункт и рекомендация.
Поддерживаемые типы: приказ и распоряжение (в т.ч. подготовленные по модулю).
"""
from __future__ import annotations

import re
from pathlib import Path

from app.approval.matrix import get_matrix
from app.templates.catalog import score_templates

OK, WARN, FAIL, INFO = "ok", "warn", "fail", "info"
ICON = {OK: "✅", WARN: "⚠️", FAIL: "❌", INFO: "ℹ️"}
DATE_ANY = re.compile(r"\b(\d{1,4})[./-](\d{1,2})[./-](\d{2,4})\b")
NAME_RE = re.compile(r"[А-ЯЁ]\.\s?[А-ЯЁ]\.\s?[А-ЯЁ][а-яё]+")


def _extract(file: Path) -> tuple[list[str], object | None]:
    if file.suffix.lower() == ".docx":
        import docx

        d = docx.Document(str(file))
        paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                seen = set()
                for c in row.cells:
                    if c._tc in seen:
                        continue
                    seen.add(c._tc)
                    if c.text.strip():
                        paras.append(c.text.strip())
        return paras, d
    if file.suffix.lower() == ".pdf":
        import pdfplumber

        with pdfplumber.open(str(file)) as pdf:
            text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        return [ln.strip() for ln in text.split("\n") if ln.strip()], None
    raise ValueError("Поддерживаются только .docx и .pdf")


def _font_stats(d) -> tuple[str | None, float | None]:
    from collections import Counter

    fonts, sizes = Counter(), Counter()
    normal = d.styles["Normal"].font
    for p in d.paragraphs:
        if len(p.text.strip()) < 40:
            continue
        for r in p.runs:
            n = len(r.text)
            f = r.font.name or (p.style.font.name if p.style is not None else None) or normal.name
            sz = r.font.size or (p.style.font.size if p.style is not None else None) or normal.size
            if f:
                fonts[f] += n
            if sz:
                sizes[sz.pt] += n
    return (fonts.most_common(1)[0][0] if fonts else None, sizes.most_common(1)[0][0] if sizes else None)


def review_document(file: Path, hint: str = "") -> dict:
    paras, d = _extract(Path(file))
    full = "\n".join(paras)
    upper = full.upper()
    items: list[dict] = []

    def add(title, status, ref, rec="", detail=""):
        items.append({"title": title, "status": status, "ref": ref, "rec": rec, "detail": detail})

    kind = "order" if re.search(r"^\s*ПРИКАЗ\s*$", full, re.M) else \
        "disposition" if re.search(r"^\s*РАСПОРЯЖЕНИЕ\s*$", full, re.M) else None
    if kind is None:
        low = (hint + " " + full[:400]).lower()
        kind = "order" if "приказ" in low else "disposition" if "распоряж" in low else None
    kind_title = {"order": "приказ", "disposition": "распоряжение"}.get(kind, "не определён")

    # 1. Вид документа и заголовок
    if kind:
        add("Указано наименование вида документа", OK, "instr_dou:8.6")
    else:
        add("Указано наименование вида документа", FAIL, "instr_dou:8.6",
            "Укажите вид документа (ПРИКАЗ / РАСПОРЯЖЕНИЕ) отдельной строкой.")
    title = next((p for p in paras if re.match(r"^(О|Об)\s", p)), None)
    if title:
        add("Есть заголовок документа («О …»)", OK, "instr_dou:8.15", detail=title[:100])
    else:
        add("Есть заголовок документа («О …»)", FAIL, "instr_dou:8.15",
            "Добавьте заголовок, кратко отражающий содержание и согласованный с видом документа.")

    # 2. Распорядительная часть
    if kind:
        word, ref = ("ПРИКАЗЫВАЮ:", "instr_dou:9.5.4") if kind == "order" else ("ОБЯЗЫВАЮ:", "instr_dou:9.7.4")
        if word in paras:
            add(f"Распорядительная часть начинается словом «{word}» отдельной строкой", OK, ref)
        elif word.rstrip(":") in upper:
            add(f"Распорядительная часть начинается словом «{word}» отдельной строкой", FAIL, ref,
                f"Напишите «{word}» прописными буквами, с двоеточием, отдельной строкой без кавычек и отступа.",
                "слово найдено, но оформлено иначе")
        else:
            add(f"Распорядительная часть начинается словом «{word}» отдельной строкой", FAIL, ref,
                f"Отделите констатирующую часть от распорядительной словом «{word}».")

    # 3. Служебные элементы модуля
    blanks = [p for p in paras if re.search(r"_{4,}", p) and "ВИЗУАЛИЗАЦИЯ" not in p]
    sign_line = [p for p in blanks if re.fullmatch(r"_+", p)]
    blanks = [p for p in blanks if p not in sign_line]
    if blanks:
        add("Нет незаполненных линий «____»", FAIL, "mr_modules:3.3",
            "Замените линии фактическими данными.", f"{len(blanks)} фрагм.: «{blanks[0][:70]}»")
    else:
        add("Нет незаполненных линий «____»", OK, "mr_modules:3.3")
    notes = re.findall(r"\[\d{1,2}\]", full)
    brackets = re.findall(r"\[[^\]\d][^\]]{0,80}\]", full)
    if notes or brackets:
        add("Удалены номера примечаний и квадратные скобки", FAIL, "mr_modules:3.3",
            "Удалите номера примечаний [N], квадратные скобки и неиспользованные варианты.",
            ", ".join((notes + brackets)[:4]))
    else:
        add("Удалены номера примечаний и квадратные скобки", OK, "mr_modules:3.3")
    variants = re.findall(r"\b\w+ / \w+\b", "\n".join(p for p in paras if "ДАТА / РЕГ" not in p))
    if variants:
        add("Выбран один вариант формулировки (нет «вариант / вариант»)", WARN, "mr_modules:3.3",
            "Оставьте нужный вариант, удалите второй и знак «/».", ", ".join(variants[:3]))
    if "МОДУЛЬНЫЙ ДОКУМЕНТ" in upper or re.search(r"^\s*УТВЕРЖДАЮ", full, re.M):
        add("Удалены гриф «УТВЕРЖДАЮ» и пометка «МОДУЛЬНЫЙ ДОКУМЕНТ»", FAIL, "mr_modules:3.3",
            "Удалите гриф утверждения вверху первой страницы и пометку «МОДУЛЬНЫЙ ДОКУМЕНТ».")
    else:
        add("Удалены гриф «УТВЕРЖДАЮ» и пометка «МОДУЛЬНЫЙ ДОКУМЕНТ»", OK, "mr_modules:3.3")
    if "ПРИМЕЧАНИЯ ПО ЗАПОЛНЕНИЮ" in upper:
        add("Удалены пояснения по заполнению", FAIL, "mr_modules:3.3", "Удалите раздел «Примечания по заполнению».")
    zones = ("ДАТА" in upper and "РЕГИСТРАЦИОННЫЙ НОМЕР" in upper, "ВИЗУАЛИЗАЦИЯ ЭП" in upper)
    if all(zones):
        add("Сохранены служебные зоны даты/номера и визуализации ЭП", OK, "mr_modules:3.3")
    else:
        add("Сохранены служебные зоны даты/номера и визуализации ЭП", WARN, "mr_modules:3.3",
            "Не удаляйте зоны для даты, регистрационного номера и визуализации электронной подписи.")

    # 4. Контроль исполнения и сроки поручений
    ctrl_ref = "instr_dou:9.5.9" if kind == "order" else "instr_dou:9.7.5"
    if re.search(r"контрол\w* за исполнением|оставляю за собой|ответственным за исполнение", full, re.I):
        add("Определено лицо, контролирующее исполнение", OK, ctrl_ref)
    else:
        add("Определено лицо, контролирующее исполнение", FAIL, ctrl_ref,
            "Добавьте пункт «Контроль за исполнением … возложить на …» или «оставляю за собой».")
    operative = full.split("ПРИКАЗЫВАЮ:" if kind == "order" else "ОБЯЗЫВАЮ:", 1)[-1]
    has_terms = bool(re.search(r"\bдо\s+«?\d|\bс\s+«?\d|\bсрок|\d{1,2}\s+(январ|феврал|март|апрел|ма[яй]|июн|июл|"
                               r"август|сентябр|октябр|ноябр|декабр)", operative, re.I))
    add("Поручения содержат сроки исполнения", OK if has_terms else WARN,
        "instr_dou:9.5.5" if kind == "order" else "instr_dou:9.7.5",
        "" if has_terms else "Каждое поручение должно определять действие, исполнителя и срок.")

    # 5. Даты
    bad_dates = []
    for m in DATE_ANY.finditer(full):
        a, b, c = m.groups()
        if not (len(a) == 2 and len(b) == 2 and len(c) == 4 and m.group(0).count(".") == 2):
            bad_dates.append(m.group(0))
    if bad_dates:
        add("Даты в формате «14.10.2026» или «14 октября 2026 г.»", WARN, "instr_dou:8.7",
            "Приведите даты к формату «ДД.ММ.ГГГГ» или «ДД месяца ГГГГ г.».", ", ".join(bad_dates[:4]))
    else:
        add("Даты в формате «14.10.2026» или «14 октября 2026 г.»", OK, "instr_dou:8.7")

    # 6. Подпись
    tail = "\n".join(paras[-12:])
    if NAME_RE.search(tail) and re.search(r"Председател|директор|руководител|заместител", tail, re.I):
        add("Реквизит «подпись»: должность, инициалы и фамилия", OK, "instr_dou:8.19.2")
    else:
        add("Реквизит «подпись»: должность, инициалы и фамилия", FAIL, "instr_dou:8.19.2",
            "Укажите должность подписанта, инициалы и фамилию (например, «А. Б. Васильев»).")

    # 7. Оформление (только .docx)
    if d is not None:
        font, size = _font_stats(d)
        if font and "times new roman" not in font.lower():
            add("Шрифт основного текста Times New Roman", FAIL, "instr_dou:8.16",
                "Используйте шрифт Times New Roman.", f"обнаружен {font}")
        elif font:
            add("Шрифт основного текста Times New Roman", OK, "instr_dou:8.16")
        if size and abs(size - 13) > 0.1:
            add("Размер шрифта основного текста 13 пт", WARN, "instr_dou:8.16",
                "Основной текст оформляется шрифтом 13 пт (12 пт допускается только во внутренней переписке "
                "объёмом более 5 листов).", f"обнаружено {size:g} пт")
        elif size:
            add("Размер шрифта основного текста 13 пт", OK, "instr_dou:8.16")
        sec = d.sections[0]
        mm = {k: round(getattr(sec, f"{k}_margin").mm) if getattr(sec, f"{k}_margin") is not None else None
              for k in ("left", "right", "top", "bottom")}
        problems = []
        if mm["left"] is not None and abs(mm["left"] - 30) > 1:
            problems.append(f"левое {mm['left']} мм (нужно 30)")
        if mm["right"] is not None and mm["right"] < 19:
            problems.append(f"правое {mm['right']} мм (не менее 20)")
        if mm["top"] is not None and mm["top"] < 24:
            problems.append(f"верхнее {mm['top']} мм (не менее 25)")
        if mm["bottom"] is not None and mm["bottom"] < 14:
            problems.append(f"нижнее {mm['bottom']} мм (не менее 15)")
        add("Поля страницы: левое 30 мм, правое ≥20, верхнее ≥25, нижнее ≥15", WARN if problems else OK,
            "instr_dou:8.16", "Исправьте поля страницы." if problems else "", "; ".join(problems))

    # 8. Пояснительная записка и визирование — подсказки
    chars = len(full)
    if chars > 5400:
        add("Объём более ~3 страниц: нужна пояснительная записка", INFO, "instr_dou:10.12",
            "Для приказов и распоряжений объёмом более трёх страниц пояснительная записка обязательна.")
    tid = None
    scored = score_templates(full[:1500])
    if scored and scored[0][0] >= 2.0:
        tid = scored[0][1]
    from app.templates.catalog import load_templates
    tpl = load_templates().get(tid) if tid else None
    if tpl and tpl.get("matrix_row") in get_matrix().rows:
        a = get_matrix().lookup(tpl["matrix_row"])
        visas = ", ".join(a.required) + ("; при условии: " + "; ".join(f"{u} ({c})" for u, c in a.conditional)
                                          if a.conditional else "")
        add(f"Обязательное визирование по Матрице: {visas}", INFO, "instr_dou:10.11",
            "Включите указанные подразделения в маршрут визирования в ИИС ДО.")

    fails = sum(i["status"] == FAIL for i in items)
    warns = sum(i["status"] == WARN for i in items)
    head = (f"**Проверка документа «{Path(file).name}»** — тип: {kind_title}"
            + (f", распознан модуль «{tpl['title']}»" if tpl else "") + "\n\n")
    head += ("**Итог: документ соответствует проверенным требованиям.**" if not fails and not warns else
             f"**Итог: замечаний — {fails}, предупреждений — {warns}.** Рекомендации по доработке ниже.")
    rows = ["| | Требование | Пункт | Рекомендация |", "|---|---|---|---|"]
    for i in sorted(items, key=lambda x: {FAIL: 0, WARN: 1, INFO: 2, OK: 3}[x["status"]]):
        doc_id, sec = i["ref"].split(":", 1)
        src = ("Инстр., п. " if doc_id == "instr_dou" else "Метод. рек., п. ") + sec
        detail = f" _({i['detail']})_" if i["detail"] else ""
        rows.append(f"| {ICON[i['status']]} | {i['title']}{detail} | {src} | {i['rec'] or '—'} |")
    md = head + "\n\n" + "\n".join(rows) + \
        "\n\n_Проверка выполнена правилами по Инструкции и методическим рекомендациям; содержательную " \
        "экспертизу выполняют визирующие подразделения (п. 10.7)._"
    return {"ok": fails == 0, "issues": fails + warns, "items": items, "markdown": md, "kind": kind}
